#!/usr/bin/env python3
"""Safe native document creation, editing, conversion, and preview CLI."""

from __future__ import annotations

import argparse
import asyncio
import base64
import binascii
import csv
import html
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable


SCHEMA_VERSION = 1
MAX_INPUT_BYTES = 250 * 1024 * 1024
MAX_ZIP_ENTRIES = 20_000
MAX_ZIP_UNCOMPRESSED = 1024 * 1024 * 1024
MAX_ZIP_RATIO = 200
MAX_PDF_PAGES = 2_000
MAX_SHEETS = 256
MAX_CELLS = 2_000_000
MAX_SLIDES = 1_000
MAX_NOTEBOOK_CELLS = 20_000
MAX_TEXT_BYTES = 4 * 1024 * 1024
MAX_NOTEBOOK_PREVIEW_BYTES = 12 * 1024 * 1024
MIN_TIMEOUT_SECONDS = 1
MAX_CONVERSION_TIMEOUT_SECONDS = 900
MAX_NOTEBOOK_TIMEOUT_SECONDS = 3_600
OFFICE_FORMATS = {".docx", ".xlsx", ".pptx"}
LEGACY_OFFICE_FORMATS = {".doc", ".xls", ".ppt"}
MACRO_FORMATS = {".docm", ".xlsm", ".pptm"}
SUPPORTED_FORMATS = OFFICE_FORMATS | {".csv", ".pdf", ".ipynb", ".html"}
PACKAGE_IMPORTS = {
    "python-docx": "docx",
    "openpyxl": "openpyxl",
    "python-pptx": "pptx",
    "pypdf": "pypdf",
    "pypdfium2": "pypdfium2",
    "reportlab": "reportlab",
    "nbformat": "nbformat",
    "nbclient": "nbclient",
    "ipykernel": "ipykernel",
    "Pillow": "PIL",
}


class WorkbenchError(RuntimeError):
    pass


def workspace_root() -> Path:
    """Return the current task workspace after resolving any directory symlink."""
    return Path.cwd().resolve(strict=True)


def resolve_workspace_path(path: Path, label: str, *, must_exist: bool = False) -> Path:
    """Resolve a caller-controlled path and keep it inside the current workspace."""
    root = workspace_root()
    candidate = path.expanduser()
    if not candidate.is_absolute():
        candidate = root / candidate
    try:
        resolved = candidate.resolve(strict=must_exist)
    except OSError as exc:
        raise WorkbenchError(f"cannot resolve {label} {path}: {exc}") from exc
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise WorkbenchError(f"{label} must stay inside the current project workspace: {path}") from exc
    return resolved


class VisibleHTMLTextParser(HTMLParser):
    """Extract visible text without executing or preserving active markup."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._hidden_depth = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "template", "noscript"}:
            self._hidden_depth += 1
        elif not self._hidden_depth and tag.lower() in {"p", "div", "section", "article", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "template", "noscript"} and self._hidden_depth:
            self._hidden_depth -= 1
        elif not self._hidden_depth and tag.lower() in {"p", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._hidden_depth:
            self.parts.append(data)

    def text(self) -> str:
        return re.sub(r"\n{3,}", "\n\n", "".join(self.parts)).strip()


def emit(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True))


def read_json(path: Path) -> dict[str, Any]:
    path = resolve_workspace_path(path, "JSON specification", must_exist=True)
    try:
        raw = path.read_text(encoding="utf-8")
    except OSError as exc:
        raise WorkbenchError(f"cannot read JSON specification {path}: {exc}") from exc
    if len(raw.encode("utf-8")) > MAX_TEXT_BYTES:
        raise WorkbenchError(f"JSON specification exceeds {MAX_TEXT_BYTES} bytes")
    try:
        value = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise WorkbenchError(f"invalid JSON specification {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise WorkbenchError("JSON specification root must be an object")
    return value


def atomic_target(output: Path) -> tuple[Path, Path]:
    output = resolve_workspace_path(output, "output")
    output.parent.mkdir(parents=True, exist_ok=True)
    suffix = f".tmp{output.suffix}" if output.suffix else ".tmp"
    handle, temporary = tempfile.mkstemp(prefix=f".{output.stem}.", suffix=suffix, dir=output.parent)
    os.close(handle)
    return output, Path(temporary)


def atomic_commit(temporary: Path, output: Path) -> None:
    os.replace(temporary, output)


def bounded_text(value: Any, label: str) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise WorkbenchError(f"{label} must be a string")
    if len(value.encode("utf-8")) > MAX_TEXT_BYTES:
        raise WorkbenchError(f"{label} exceeds {MAX_TEXT_BYTES} bytes")
    return value


def bounded_integer(value: Any, label: str, minimum: int, maximum: int) -> int:
    if isinstance(value, bool):
        raise WorkbenchError(f"{label} must be an integer")
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise WorkbenchError(f"{label} must be an integer") from exc
    if parsed < minimum or parsed > maximum:
        raise WorkbenchError(f"{label} must be between {minimum} and {maximum}")
    return parsed


def notebook_text(value: Any, label: str) -> str:
    """Normalize the two text encodings permitted by the notebook v4 schema."""
    if isinstance(value, str):
        return bounded_text(value, label)
    if isinstance(value, list) and all(isinstance(part, str) for part in value):
        return bounded_text("".join(value), label)
    raise WorkbenchError(f"{label} must be a string or an array of strings")


def notebook_cell(cell_type: str, source: str, tags: list[str], identifier: str) -> dict[str, Any]:
    cell: dict[str, Any] = {
        "cell_type": cell_type,
        "id": identifier,
        "metadata": {"tags": tags},
        "source": source,
    }
    if cell_type == "code":
        cell.update({"execution_count": None, "outputs": []})
    return cell


def next_notebook_cell_id(cells: list[dict[str, Any]]) -> str:
    existing = {str(cell.get("id", "")) for cell in cells}
    sequence = len(cells) + 1
    while f"cell-{sequence:04d}" in existing:
        sequence += 1
    return f"cell-{sequence:04d}"


def load_notebook(path: Path) -> dict[str, Any]:
    """Load and bound a notebook without requiring Jupyter to preview native files."""
    try:
        raw = path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        raise WorkbenchError(f"cannot read notebook {path}: {exc}") from exc
    try:
        notebook = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise WorkbenchError(f"invalid notebook JSON: {exc}") from exc
    if not isinstance(notebook, dict):
        raise WorkbenchError("notebook root must be an object")
    if notebook.get("nbformat") != 4:
        raise WorkbenchError("only notebook format version 4 is supported")
    minor = notebook.get("nbformat_minor")
    if not isinstance(minor, int) or isinstance(minor, bool) or minor < 0:
        raise WorkbenchError("notebook nbformat_minor must be a non-negative integer")
    metadata = notebook.get("metadata")
    if not isinstance(metadata, dict):
        raise WorkbenchError("notebook metadata must be an object")
    cells = notebook.get("cells")
    if not isinstance(cells, list) or len(cells) > MAX_NOTEBOOK_CELLS:
        raise WorkbenchError(f"notebook cells must be an array with at most {MAX_NOTEBOOK_CELLS} items")

    content_bytes = 0
    for index, cell in enumerate(cells):
        if not isinstance(cell, dict):
            raise WorkbenchError(f"notebook cell {index + 1} must be an object")
        cell_type = cell.get("cell_type")
        if cell_type not in {"markdown", "code", "raw"}:
            raise WorkbenchError(f"notebook cell {index + 1} has unsupported type: {cell_type!r}")
        cell_metadata = cell.get("metadata")
        if not isinstance(cell_metadata, dict):
            raise WorkbenchError(f"notebook cell {index + 1} metadata must be an object")
        identifier = cell.get("id")
        if identifier is not None and (not isinstance(identifier, str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,64}", identifier)):
            raise WorkbenchError(f"notebook cell {index + 1} has an invalid id")
        source = notebook_text(cell.get("source", ""), f"notebook cell {index + 1} source")
        cell["source"] = source
        content_bytes += len(source.encode("utf-8"))
        if content_bytes > MAX_NOTEBOOK_PREVIEW_BYTES:
            raise WorkbenchError(f"notebook text and output content exceeds {MAX_NOTEBOOK_PREVIEW_BYTES} bytes")
        if cell_type != "code":
            attachments = cell.get("attachments")
            if attachments is not None and not isinstance(attachments, dict):
                raise WorkbenchError(f"notebook cell {index + 1} attachments must be an object")
            continue

        execution_count = cell.get("execution_count")
        if execution_count is not None and (not isinstance(execution_count, int) or isinstance(execution_count, bool) or execution_count < 0):
            raise WorkbenchError(f"notebook cell {index + 1} execution_count must be null or a non-negative integer")
        outputs = cell.get("outputs")
        if not isinstance(outputs, list):
            raise WorkbenchError(f"notebook cell {index + 1} outputs must be an array")
        for output_index, output in enumerate(outputs):
            if not isinstance(output, dict):
                raise WorkbenchError(f"notebook cell {index + 1} output {output_index + 1} must be an object")
            output_type = output.get("output_type")
            if output_type not in {"stream", "display_data", "execute_result", "error"}:
                raise WorkbenchError(f"notebook cell {index + 1} output {output_index + 1} has an unsupported type")
            if output_type == "stream":
                stream = notebook_text(output.get("text", ""), "notebook stream output")
                output["text"] = stream
                content_bytes += len(stream.encode("utf-8"))
            elif output_type == "error":
                traceback = output.get("traceback", [])
                if not isinstance(traceback, list) or not all(isinstance(line, str) for line in traceback):
                    raise WorkbenchError("notebook error traceback must be an array of strings")
                content_bytes += sum(len(line.encode("utf-8")) for line in traceback)
            else:
                data = output.get("data")
                if not isinstance(data, dict):
                    raise WorkbenchError("notebook rich output data must be an object")
                plain = data.get("text/plain")
                if plain is not None:
                    plain_text = notebook_text(plain, "notebook text output")
                    data["text/plain"] = plain_text
                    content_bytes += len(plain_text.encode("utf-8"))
                image_data = data.get("image/png")
                if image_data is not None and not isinstance(image_data, str):
                    raise WorkbenchError("notebook PNG output must be a base64 string")
                if isinstance(image_data, str):
                    content_bytes += len(image_data.encode("ascii", errors="ignore"))
            if content_bytes > MAX_NOTEBOOK_PREVIEW_BYTES:
                raise WorkbenchError(f"notebook text and output content exceeds {MAX_NOTEBOOK_PREVIEW_BYTES} bytes")
    return notebook


def write_notebook(notebook: dict[str, Any], output: Path) -> None:
    target, temporary = atomic_target(output)
    try:
        temporary.write_text(json.dumps(notebook, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        atomic_commit(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def ensure_keys(spec: dict[str, Any], allowed: set[str], label: str) -> None:
    unknown = sorted(set(spec) - allowed)
    if unknown:
        raise WorkbenchError(f"{label} contains unsupported fields: {', '.join(unknown)}")


def ensure_input(path: Path, *, allow_legacy: bool = False) -> Path:
    path = resolve_workspace_path(path, "input", must_exist=True)
    if not path.is_file():
        raise WorkbenchError(f"input file not found: {path}")
    size = path.stat().st_size
    if size > MAX_INPUT_BYTES:
        raise WorkbenchError(f"input exceeds {MAX_INPUT_BYTES} bytes: {size}")
    suffix = path.suffix.lower()
    if suffix in MACRO_FORMATS:
        raise WorkbenchError(f"macro-enabled Office files are rejected: {suffix}")
    allowed = SUPPORTED_FORMATS | (LEGACY_OFFICE_FORMATS if allow_legacy else set())
    if suffix not in allowed:
        raise WorkbenchError(f"unsupported input format: {suffix or '<none>'}")
    return path


def validate_zip_package(path: Path) -> dict[str, int]:
    try:
        archive = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile) as exc:
        raise WorkbenchError(f"invalid OOXML/Notebook ZIP package: {exc}") from exc
    compressed = 0
    uncompressed = 0
    with archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_ENTRIES:
            raise WorkbenchError(f"package has {len(entries)} entries; limit is {MAX_ZIP_ENTRIES}")
        for entry in entries:
            normalized = entry.filename.replace("\\", "/")
            parts = [part for part in normalized.split("/") if part not in ("", ".")]
            if normalized.startswith("/") or ".." in parts:
                raise WorkbenchError(f"unsafe package entry path: {entry.filename}")
            if entry.file_size < 0 or entry.compress_size < 0:
                raise WorkbenchError(f"invalid package entry size: {entry.filename}")
            compressed += entry.compress_size
            uncompressed += entry.file_size
            if uncompressed > MAX_ZIP_UNCOMPRESSED:
                raise WorkbenchError(f"package expands beyond {MAX_ZIP_UNCOMPRESSED} bytes")
        ratio = uncompressed / max(compressed, 1)
        if ratio > MAX_ZIP_RATIO:
            raise WorkbenchError(f"package expansion ratio {ratio:.1f}:1 exceeds {MAX_ZIP_RATIO}:1")
    return {"entries": len(entries), "compressedBytes": compressed, "uncompressedBytes": uncompressed}


def package_status() -> dict[str, bool]:
    return {name: importlib.util.find_spec(module) is not None for name, module in PACKAGE_IMPORTS.items()}


def find_soffice() -> str | None:
    candidates = [
        os.environ.get("LIBREOFFICE_BIN", ""),
        shutil.which("soffice") or "",
        shutil.which("libreoffice") or "",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(Path(candidate).resolve())
    return None


def require_module(distribution: str, module: str | None = None) -> None:
    module = module or PACKAGE_IMPORTS.get(distribution, distribution.replace("-", "_"))
    if importlib.util.find_spec(module) is None:
        raise WorkbenchError(f"missing Python package {distribution}; use requirements.lock in a managed environment")


def command_doctor(_: argparse.Namespace) -> dict[str, Any]:
    packages = package_status()
    soffice = find_soffice()
    required_packages = ("python-docx", "openpyxl", "python-pptx", "pypdf", "pypdfium2", "reportlab", "Pillow")
    optional_packages = ("nbformat", "nbclient", "ipykernel")
    core_ready = all(packages.get(name, False) for name in required_packages)
    return {
        "ok": core_ready,
        "command": "doctor",
        "python": sys.executable,
        "pythonVersion": sys.version.split()[0],
        "packages": packages,
        "packageTiers": {"required": list(required_packages), "optionalNotebookExecution": list(optional_packages)},
        "libreOffice": {"available": soffice is not None, "path": soffice},
        "capabilities": {
            "nativeCreate": all(packages.get(name, False) for name in ("python-docx", "openpyxl", "python-pptx", "reportlab")),
            "nativeNotebookPreview": True,
            "pdfRender": packages.get("pypdfium2", False),
            "notebookExecution": packages.get("nbclient", False) and packages.get("nbformat", False) and packages.get("ipykernel", False),
            "officePdfConversion": soffice is not None,
        },
    }


def create_docx(spec: dict[str, Any], output: Path) -> None:
    require_module("python-docx", "docx")
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    ensure_keys(spec, {"title", "subtitle", "author", "sections"}, "DOCX specification")
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    title = bounded_text(spec.get("title", "Untitled document"), "title")
    subtitle = bounded_text(spec.get("subtitle", ""), "subtitle")
    author = bounded_text(spec.get("author", "Synon Biomed"), "author")
    document.core_properties.title = title
    document.core_properties.author = author
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Title"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(title)
    if subtitle:
        sub = document.add_paragraph(subtitle)
        sub.style = document.styles["Subtitle"]
    sections = spec.get("sections", [])
    if not isinstance(sections, list):
        raise WorkbenchError("sections must be an array")
    for index, item in enumerate(sections):
        if not isinstance(item, dict):
            raise WorkbenchError(f"sections[{index}] must be an object")
        ensure_keys(item, {"heading", "level", "blocks"}, f"sections[{index}]")
        heading = bounded_text(item.get("heading", ""), f"sections[{index}].heading")
        level = int(item.get("level", 1))
        if heading:
            document.add_heading(heading, level=max(1, min(9, level)))
        blocks = item.get("blocks", [])
        if not isinstance(blocks, list):
            raise WorkbenchError(f"sections[{index}].blocks must be an array")
        for block_index, block in enumerate(blocks):
            if not isinstance(block, dict):
                raise WorkbenchError(f"block {block_index} must be an object")
            kind = block.get("type")
            if kind == "paragraph":
                ensure_keys(block, {"type", "text"}, "paragraph block")
                document.add_paragraph(bounded_text(block.get("text", ""), "paragraph text"))
            elif kind == "bullets":
                ensure_keys(block, {"type", "items"}, "bullets block")
                items = block.get("items", [])
                if not isinstance(items, list):
                    raise WorkbenchError("bullet items must be an array")
                for bullet in items:
                    document.add_paragraph(bounded_text(bullet, "bullet"), style="List Bullet")
            elif kind == "table":
                ensure_keys(block, {"type", "headers", "rows"}, "table block")
                headers = block.get("headers", [])
                rows = block.get("rows", [])
                if not isinstance(headers, list) or not isinstance(rows, list) or not headers:
                    raise WorkbenchError("table requires non-empty headers and rows array")
                table = document.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for cell, value in zip(table.rows[0].cells, headers):
                    cell.text = bounded_text(str(value), "table header")
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                for row in rows:
                    if not isinstance(row, list) or len(row) != len(headers):
                        raise WorkbenchError("every table row must match the header width")
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = bounded_text("" if value is None else str(value), "table cell")
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            elif kind == "image":
                ensure_keys(block, {"type", "path", "caption", "widthInches"}, "image block")
                image_path = resolve_workspace_path(
                    Path(bounded_text(block.get("path"), "image path")), "image path", must_exist=True
                )
                if not image_path.is_file():
                    raise WorkbenchError(f"image not found: {image_path}")
                width = float(block.get("widthInches", 5.8))
                document.add_picture(str(image_path), width=Inches(max(0.5, min(7.0, width))))
                caption = bounded_text(block.get("caption", ""), "image caption")
                if caption:
                    p = document.add_paragraph(caption)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                raise WorkbenchError(f"unsupported DOCX block type: {kind!r}")
    output, temporary = atomic_target(output)
    try:
        document.save(temporary)
        atomic_commit(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def create_xlsx(spec: dict[str, Any], output: Path) -> None:
    require_module("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ensure_keys(spec, {"title", "allowFormulas", "sheets"}, "XLSX specification")
    sheets = spec.get("sheets", [])
    if not isinstance(sheets, list) or not sheets:
        raise WorkbenchError("XLSX specification requires at least one sheet")
    if len(sheets) > MAX_SHEETS:
        raise WorkbenchError(f"sheet count exceeds {MAX_SHEETS}")
    workbook = Workbook()
    workbook.remove(workbook.active)
    allow_formulas = bool(spec.get("allowFormulas", False))
    total_cells = 0
    for sheet_index, sheet_spec in enumerate(sheets):
        if not isinstance(sheet_spec, dict):
            raise WorkbenchError(f"sheets[{sheet_index}] must be an object")
        ensure_keys(sheet_spec, {"name", "headers", "rows", "freeze", "autoFilter", "widths", "numberFormats"}, f"sheets[{sheet_index}]")
        name = bounded_text(sheet_spec.get("name", f"Sheet {sheet_index + 1}"), "sheet name")[:31]
        if not name or any(character in name for character in "[]:*?/\\"):
            raise WorkbenchError(f"invalid sheet name: {name!r}")
        worksheet = workbook.create_sheet(name)
        headers = sheet_spec.get("headers", [])
        rows = sheet_spec.get("rows", [])
        if not isinstance(headers, list) or not isinstance(rows, list):
            raise WorkbenchError("sheet headers and rows must be arrays")
        width = len(headers) if headers else max((len(row) for row in rows if isinstance(row, list)), default=0)
        if width == 0:
            raise WorkbenchError(f"sheet {name!r} has no columns")
        if headers:
            worksheet.append([safe_cell_value(value, allow_formulas) for value in headers])
        for row in rows:
            if not isinstance(row, list) or len(row) > width:
                raise WorkbenchError(f"row in sheet {name!r} is not an array within width {width}")
            worksheet.append([safe_cell_value(value, allow_formulas) for value in row])
        total_cells += worksheet.max_row * worksheet.max_column
        if total_cells > MAX_CELLS:
            raise WorkbenchError(f"workbook cell count exceeds {MAX_CELLS}")
        if headers:
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2B2B2B")
                cell.alignment = Alignment(vertical="center")
            worksheet.row_dimensions[1].height = 22
        freeze = sheet_spec.get("freeze", "A2" if headers else None)
        worksheet.freeze_panes = freeze
        if sheet_spec.get("autoFilter", bool(headers)) and headers:
            worksheet.auto_filter.ref = worksheet.dimensions
        widths = sheet_spec.get("widths", {})
        if not isinstance(widths, dict):
            raise WorkbenchError("widths must be an object")
        for column_index in range(1, worksheet.max_column + 1):
            letter = get_column_letter(column_index)
            measured = max((len(str(worksheet.cell(row=row, column=column_index).value or "")) for row in range(1, worksheet.max_row + 1)), default=8)
            worksheet.column_dimensions[letter].width = max(8.0, min(60.0, float(widths.get(letter, measured + 2))))
        formats = sheet_spec.get("numberFormats", {})
        if not isinstance(formats, dict):
            raise WorkbenchError("numberFormats must be an object")
        for letter, number_format in formats.items():
            for cell in worksheet[str(letter)]:
                if cell.row > (1 if headers else 0):
                    cell.number_format = bounded_text(number_format, "number format")
        worksheet.sheet_view.showGridLines = False
    workbook.properties.title = bounded_text(spec.get("title", "Workbook"), "title")
    output, temporary = atomic_target(output)
    try:
        workbook.save(temporary)
        atomic_commit(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def safe_cell_value(value: Any, allow_formulas: bool) -> Any:
    if isinstance(value, str):
        if value.startswith("=") and not allow_formulas:
            return "'" + value
        if value.startswith(("+", "-", "@")) and not allow_formulas:
            return "'" + value
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    raise WorkbenchError(f"unsupported spreadsheet cell value: {type(value).__name__}")


def create_pptx(spec: dict[str, Any], output: Path) -> None:
    require_module("python-pptx", "pptx")
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    ensure_keys(spec, {"title", "author", "slides"}, "PPTX specification")
    slides = spec.get("slides", [])
    if not isinstance(slides, list) or not slides:
        raise WorkbenchError("PPTX specification requires at least one slide")
    if len(slides) > MAX_SLIDES:
        raise WorkbenchError(f"slide count exceeds {MAX_SLIDES}")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = bounded_text(spec.get("title", "Presentation"), "title")
    presentation.core_properties.author = bounded_text(spec.get("author", "Synon Biomed"), "author")
    for index, slide_spec in enumerate(slides):
        if not isinstance(slide_spec, dict):
            raise WorkbenchError(f"slides[{index}] must be an object")
        ensure_keys(slide_spec, {"layout", "title", "subtitle", "bullets", "image", "caption"}, f"slides[{index}]")
        layout = slide_spec.get("layout", "content")
        layout_index = 0 if layout == "title" else 1
        slide = presentation.slides.add_slide(presentation.slide_layouts[layout_index])
        title = bounded_text(slide_spec.get("title", ""), "slide title")
        if slide.shapes.title:
            slide.shapes.title.text = title
            slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
            slide.shapes.title.text_frame.paragraphs[0].font.bold = True
        if layout == "title":
            subtitle = bounded_text(slide_spec.get("subtitle", ""), "slide subtitle")
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = subtitle
        elif layout == "image":
            image_path = resolve_workspace_path(
                Path(bounded_text(slide_spec.get("image"), "slide image")), "slide image", must_exist=True
            )
            if not image_path.is_file():
                raise WorkbenchError(f"slide image not found: {image_path}")
            require_module("Pillow", "PIL")
            from PIL import Image as PillowImage

            with PillowImage.open(image_path) as source_image:
                image_width, image_height = source_image.size
            if image_width <= 0 or image_height <= 0:
                raise WorkbenchError(f"slide image has invalid dimensions: {image_path}")
            box_width, box_height = 10.9, 4.9
            scale = min(box_width / image_width, box_height / image_height)
            width_inches, height_inches = image_width * scale, image_height * scale
            left_inches = 1.2 + (box_width - width_inches) / 2
            top_inches = 1.45 + (box_height - height_inches) / 2
            slide.shapes.add_picture(
                str(image_path),
                Inches(left_inches),
                Inches(top_inches),
                width=Inches(width_inches),
                height=Inches(height_inches),
            )
            caption = bounded_text(slide_spec.get("caption", ""), "slide caption")
            if caption:
                box = slide.shapes.add_textbox(Inches(1.2), Inches(6.45), Inches(10.9), Inches(0.45))
                paragraph = box.text_frame.paragraphs[0]
                paragraph.text = caption
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.font.size = Pt(12)
        else:
            bullets = slide_spec.get("bullets", [])
            if not isinstance(bullets, list):
                raise WorkbenchError("slide bullets must be an array")
            body = next((placeholder for placeholder in slide.placeholders if getattr(placeholder, "has_text_frame", False) and placeholder != slide.shapes.title), None)
            if body is None:
                body = slide.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(11.1), Inches(4.8))
            frame = body.text_frame
            frame.clear()
            for bullet_index, bullet in enumerate(bullets):
                paragraph = frame.paragraphs[0] if bullet_index == 0 else frame.add_paragraph()
                paragraph.text = bounded_text(bullet, "slide bullet")
                paragraph.level = 0
                paragraph.font.size = Pt(20)
                paragraph.space_after = Pt(12)
    output, temporary = atomic_target(output)
    try:
        presentation.save(temporary)
        atomic_commit(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def create_pdf(spec: dict[str, Any], output: Path) -> None:
    require_module("reportlab")
    from reportlab.lib import colors
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.pdfmetrics import registerFont
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ensure_keys(spec, {"title", "subtitle", "author", "sections"}, "PDF specification")
    output, temporary = atomic_target(output)
    registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="TitleSafe", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=28, spaceAfter=12))
    styles.add(ParagraphStyle(name="SubtitleSafe", parent=styles["Heading2"], fontName="Helvetica", fontSize=13, leading=18, textColor=colors.HexColor("#505050")))
    styles.add(ParagraphStyle(name="Heading1Safe", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=21, spaceBefore=12, spaceAfter=8))
    styles.add(ParagraphStyle(name="Heading2Safe", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=18, spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle(name="Heading3Safe", parent=styles["Heading3"], fontName="Helvetica-Bold", fontSize=11, leading=16, spaceBefore=8, spaceAfter=5))
    styles.add(ParagraphStyle(name="BodySafe", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=14, spaceAfter=8))
    styles.add(ParagraphStyle(name="TableHeaderSafe", parent=styles["BodySafe"], fontName="Helvetica-Bold", textColor=colors.white))
    styles.add(ParagraphStyle(name="CaptionSafe", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, leading=12, textColor=colors.HexColor("#666666")))

    def pdf_text(value: Any, label: str) -> str:
        escaped = html.escape(bounded_text(value, label))
        return re.sub(
            r"([\u3000-\u303f\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff\uff00-\uffef]+)",
            lambda match: f'<font name="STSong-Light">{match.group(1)}</font>',
            escaped,
        )

    story: list[Any] = []
    title = bounded_text(spec.get("title", "Untitled document"), "title")
    story.append(Paragraph(pdf_text(title, "title"), styles["TitleSafe"]))
    subtitle = bounded_text(spec.get("subtitle", ""), "subtitle")
    if subtitle:
        story.append(Paragraph(pdf_text(subtitle, "subtitle"), styles["SubtitleSafe"]))
    story.append(Spacer(1, 0.12 * inch))
    sections = spec.get("sections", [])
    if not isinstance(sections, list):
        raise WorkbenchError("sections must be an array")
    for item in sections:
        heading = bounded_text(item.get("heading", ""), "section heading")
        level = max(1, min(3, int(item.get("level", 1))))
        if heading:
            story.append(Paragraph(pdf_text(heading, "section heading"), styles[f"Heading{level}Safe"]))
        for block in item.get("blocks", []):
            kind = block.get("type")
            if kind == "paragraph":
                story.append(Paragraph(pdf_text(block.get("text", ""), "paragraph text").replace("\n", "<br/>"), styles["BodySafe"]))
            elif kind == "bullets":
                story.append(ListFlowable([ListItem(Paragraph(pdf_text(value, "bullet"), styles["BodySafe"])) for value in block.get("items", [])], bulletType="bullet", bulletFontName="Helvetica"))
            elif kind == "table":
                headers = block.get("headers", [])
                rows = block.get("rows", [])
                data = [[Paragraph(pdf_text(value, "table header"), styles["TableHeaderSafe"]) for value in headers]]
                data.extend([[Paragraph(pdf_text("" if value is None else value, "table cell"), styles["BodySafe"]) for value in row] for row in rows])
                table = Table(data, repeatRows=1)
                table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2b2b2b")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "STSong-Light"), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d0d0d0")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
                story.extend([table, Spacer(1, 0.12 * inch)])
            elif kind == "image":
                image_path = resolve_workspace_path(
                    Path(bounded_text(block.get("path"), "image path")), "image path", must_exist=True
                )
                if not image_path.is_file():
                    raise WorkbenchError(f"image not found: {image_path}")
                width = max(0.5, min(7.0, float(block.get("widthInches", 5.8)))) * inch
                image = Image(str(image_path))
                if image.imageWidth <= 0 or image.imageHeight <= 0:
                    raise WorkbenchError(f"PDF image has invalid dimensions: {image_path}")
                image.drawWidth = width
                image.drawHeight = width * image.imageHeight / image.imageWidth
                image.hAlign = "CENTER"
                story.append(image)
                caption = bounded_text(block.get("caption", ""), "image caption")
                if caption:
                    story.append(Paragraph(pdf_text(caption, "image caption"), styles["CaptionSafe"]))
            else:
                raise WorkbenchError(f"unsupported PDF block type: {kind!r}")
    document = SimpleDocTemplate(str(temporary), pagesize=A4, title=title, author=bounded_text(spec.get("author", "Synon Biomed"), "author"), rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
    try:
        document.build(story)
        atomic_commit(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def create_notebook(spec: dict[str, Any], output: Path) -> None:
    ensure_keys(spec, {"title", "kernel", "cells"}, "IPYNB specification")
    cells = spec.get("cells", [])
    if not isinstance(cells, list) or len(cells) > MAX_NOTEBOOK_CELLS:
        raise WorkbenchError(f"cells must be an array with at most {MAX_NOTEBOOK_CELLS} items")
    kernel = spec.get("kernel", {})
    if not isinstance(kernel, dict):
        raise WorkbenchError("kernel must be an object")
    ensure_keys(kernel, {"name", "displayName", "language"}, "kernel")
    kernel_spec = {"name": bounded_text(kernel.get("name", "python3"), "kernel name"), "display_name": bounded_text(kernel.get("displayName", "Python 3"), "kernel displayName"), "language": bounded_text(kernel.get("language", "python"), "kernel language")}
    notebook: dict[str, Any] = {
        "cells": [],
        "metadata": {
            "title": bounded_text(spec.get("title", "Notebook"), "title"),
            "kernelspec": kernel_spec,
            "language_info": {"name": kernel_spec["language"]},
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    for index, cell_spec in enumerate(cells):
        if not isinstance(cell_spec, dict):
            raise WorkbenchError(f"cells[{index}] must be an object")
        ensure_keys(cell_spec, {"type", "source", "tags"}, f"cells[{index}]")
        cell_type = cell_spec.get("type")
        source = bounded_text(cell_spec.get("source", ""), f"cells[{index}].source")
        tags = cell_spec.get("tags", [])
        if not isinstance(tags, list) or not all(isinstance(tag, str) for tag in tags):
            raise WorkbenchError(f"cells[{index}].tags must be an array of strings")
        if cell_type not in {"markdown", "code", "raw"}:
            raise WorkbenchError(f"unsupported notebook cell type: {cell_type!r}")
        notebook["cells"].append(notebook_cell(cell_type, source, tags, f"cell-{index + 1:04d}"))
    write_notebook(notebook, output)


def create_html(spec: dict[str, Any], output: Path) -> None:
    """Create a self-contained, passive HTML report from the document specification."""
    ensure_keys(spec, {"title", "subtitle", "author", "sections"}, "HTML specification")
    title = bounded_text(spec.get("title", "Untitled report"), "title")
    subtitle = bounded_text(spec.get("subtitle", ""), "subtitle")
    sections = spec.get("sections", [])
    if not isinstance(sections, list):
        raise WorkbenchError("sections must be an array")

    body: list[str] = []
    if subtitle:
        body.append(f'<p class="subtitle">{html.escape(subtitle)}</p>')
    for section_index, item in enumerate(sections):
        if not isinstance(item, dict):
            raise WorkbenchError(f"sections[{section_index}] must be an object")
        ensure_keys(item, {"heading", "level", "blocks"}, f"sections[{section_index}]")
        heading = bounded_text(item.get("heading", ""), f"sections[{section_index}].heading")
        level = max(1, min(3, int(item.get("level", 1))))
        if heading:
            body.append(f"<h{level + 1}>{html.escape(heading)}</h{level + 1}>")
        blocks = item.get("blocks", [])
        if not isinstance(blocks, list):
            raise WorkbenchError(f"sections[{section_index}].blocks must be an array")
        for block_index, block in enumerate(blocks):
            if not isinstance(block, dict):
                raise WorkbenchError(f"sections[{section_index}].blocks[{block_index}] must be an object")
            kind = block.get("type")
            if kind == "paragraph":
                ensure_keys(block, {"type", "text"}, "paragraph block")
                text = bounded_text(block.get("text", ""), "paragraph text")
                body.append(f"<p>{html.escape(text).replace(chr(10), '<br>')}</p>")
            elif kind == "bullets":
                ensure_keys(block, {"type", "items"}, "bullets block")
                items = block.get("items", [])
                if not isinstance(items, list):
                    raise WorkbenchError("bullet items must be an array")
                body.append("<ul>" + "".join(f"<li>{html.escape(bounded_text(value, 'bullet'))}</li>" for value in items) + "</ul>")
            elif kind == "table":
                ensure_keys(block, {"type", "headers", "rows"}, "table block")
                headers = block.get("headers", [])
                rows = block.get("rows", [])
                if not isinstance(headers, list) or not isinstance(rows, list) or not headers:
                    raise WorkbenchError("table requires non-empty headers and rows array")
                table = ["<table><thead><tr>"]
                table.extend(f"<th>{html.escape(bounded_text(str(value), 'table header'))}</th>" for value in headers)
                table.append("</tr></thead><tbody>")
                for row in rows:
                    if not isinstance(row, list) or len(row) != len(headers):
                        raise WorkbenchError("every table row must match the header width")
                    table.append("<tr>" + "".join(f"<td>{html.escape('' if value is None else str(value))}</td>" for value in row) + "</tr>")
                table.append("</tbody></table>")
                body.append("".join(table))
            else:
                raise WorkbenchError(f"unsupported HTML block type: {kind!r}")

    document = escape_document_html(title, "".join(body))
    output, temporary = atomic_target(output)
    try:
        temporary.write_text(document, encoding="utf-8")
        atomic_commit(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def command_create(args: argparse.Namespace) -> dict[str, Any]:
    spec = read_json(Path(args.spec))
    output = resolve_workspace_path(Path(args.output), "output")
    format_name = args.format.lower().lstrip(".")
    expected = "." + format_name
    if output.suffix.lower() != expected:
        raise WorkbenchError(f"output suffix {output.suffix!r} does not match --format {format_name}")
    creators = {
        "docx": create_docx,
        "xlsx": create_xlsx,
        "pptx": create_pptx,
        "pdf": create_pdf,
        "ipynb": create_notebook,
        "html": create_html,
    }
    creator = creators.get(format_name)
    if creator is None:
        raise WorkbenchError(f"unsupported create format: {format_name}")
    creator(spec, output)
    report = validate_file(output)
    return {"ok": True, "command": "create", "output": str(output.resolve()), "validation": report}


def validate_file(path: Path) -> dict[str, Any]:
    path = ensure_input(path)
    suffix = path.suffix.lower()
    report: dict[str, Any] = {"path": str(path), "format": suffix.lstrip("."), "bytes": path.stat().st_size, "valid": True}
    if suffix in OFFICE_FORMATS:
        report["package"] = validate_zip_package(path)
    if suffix == ".docx":
        require_module("python-docx", "docx")
        from docx import Document
        Document(path)
    elif suffix == ".xlsx":
        require_module("openpyxl")
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if len(workbook.sheetnames) > MAX_SHEETS:
                raise WorkbenchError(f"sheet count exceeds {MAX_SHEETS}")
            cells = sum(worksheet.max_row * worksheet.max_column for worksheet in workbook.worksheets)
            if cells > MAX_CELLS:
                raise WorkbenchError(f"workbook cell count exceeds {MAX_CELLS}")
            report.update({"sheets": len(workbook.sheetnames), "cells": cells})
        finally:
            workbook.close()
    elif suffix == ".pptx":
        require_module("python-pptx", "pptx")
        from pptx import Presentation
        presentation = Presentation(path)
        if len(presentation.slides) > MAX_SLIDES:
            raise WorkbenchError(f"slide count exceeds {MAX_SLIDES}")
        report["slides"] = len(presentation.slides)
    elif suffix == ".pdf":
        require_module("pypdf")
        from pypdf import PdfReader
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise WorkbenchError("encrypted PDF is blocked until a password is provided through an approved secret flow")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise WorkbenchError(f"PDF page count exceeds {MAX_PDF_PAGES}")
        report["pages"] = len(reader.pages)
    elif suffix == ".ipynb":
        notebook = load_notebook(path)
        report["cells"] = len(notebook["cells"])
    elif suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = sum(1 for _ in csv.reader(handle))
        report["rows"] = rows
    elif suffix == ".html":
        source = path.read_text(encoding="utf-8")
        active_patterns = (
            r"<(script|iframe|object|embed|svg|math|base|link)\b",
            r"\son[a-z]+\s*=",
            r"(?:href|src|action)\s*=\s*['\"]\s*(?:javascript:|data:text/html|https?://|//)",
            r"<meta\b[^>]*http-equiv\s*=\s*['\"]?refresh",
        )
        if any(re.search(pattern, source, re.IGNORECASE) for pattern in active_patterns):
            raise WorkbenchError("HTML contains active or externally loaded content and is not a safe preview")
    return report


def command_validate(args: argparse.Namespace) -> dict[str, Any]:
    return {"ok": True, "command": "validate", "validation": validate_file(Path(args.input))}


def inspect_file(path: Path) -> dict[str, Any]:
    validation = validate_file(path)
    path = Path(validation["path"])
    suffix = path.suffix.lower()
    result: dict[str, Any] = {"schemaVersion": SCHEMA_VERSION, "validation": validation}
    if suffix == ".docx":
        from docx import Document
        document = Document(path)
        result["document"] = {"title": document.core_properties.title or "", "author": document.core_properties.author or "", "paragraphs": [paragraph.text for paragraph in document.paragraphs], "tables": [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]}
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            result["workbook"] = {"sheets": [{"name": worksheet.title, "rows": worksheet.max_row, "columns": worksheet.max_column, "sample": [[cell.value for cell in row] for row in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 20), values_only=False)]} for worksheet in workbook.worksheets]}
        finally:
            workbook.close()
    elif suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        result["presentation"] = {"slides": [{"number": index + 1, "texts": [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text]} for index, slide in enumerate(presentation.slides)]}
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        result["pdf"] = {"metadata": {str(key): str(value) for key, value in (reader.metadata or {}).items()}, "pages": [{"number": index + 1, "text": (page.extract_text() or "")[:8000]} for index, page in enumerate(reader.pages[:50])]}
    elif suffix == ".ipynb":
        notebook = load_notebook(path)
        result["notebook"] = {"metadata": notebook["metadata"], "cells": [{"index": index, "type": cell["cell_type"], "source": cell["source"], "outputs": len(cell.get("outputs", []))} for index, cell in enumerate(notebook["cells"])]}
    elif suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            result["csv"] = {"rows": [row for _, row in zip(range(100), csv.reader(handle))]}
    elif suffix == ".html":
        result["html"] = {"text": re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", path.read_text(encoding="utf-8")))[:16000]}
    return result


def command_inspect(args: argparse.Namespace) -> dict[str, Any]:
    result = inspect_file(Path(args.input))
    if args.output:
        output, temporary = atomic_target(resolve_workspace_path(Path(args.output), "output"))
        try:
            temporary.write_text(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
            atomic_commit(temporary, output)
        finally:
            temporary.unlink(missing_ok=True)
        return {"ok": True, "command": "inspect", "output": str(output), "summary": result["validation"]}
    return {"ok": True, "command": "inspect", **result}


def replace_text_in_runs(paragraphs: Iterable[Any], find: str, replace: str, count: int) -> int:
    changed = 0
    remaining = count
    for paragraph in paragraphs:
        if find not in paragraph.text or (count > 0 and remaining <= 0):
            continue
        original = paragraph.text
        limit = remaining if count > 0 else -1
        updated = original.replace(find, replace, limit)
        replacements = original.count(find) if limit < 0 else min(original.count(find), limit)
        if updated != original:
            if paragraph.runs:
                paragraph.runs[0].text = updated
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = updated
            changed += replacements
            if count > 0:
                remaining -= replacements
    return changed


def edit_docx(path: Path, spec: dict[str, Any], output: Path) -> int:
    from docx import Document
    document = Document(path)
    changes = 0
    for operation in spec["operations"]:
        op = operation.get("op")
        if op == "replace_text":
            find = bounded_text(operation.get("find"), "find")
            if not find:
                raise WorkbenchError("replace_text.find cannot be empty")
            replace = bounded_text(operation.get("replace", ""), "replace")
            count = int(operation.get("count", 0))
            paragraphs = list(document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            changes += replace_text_in_runs(paragraphs, find, replace, count)
        elif op == "append_section":
            document.add_heading(bounded_text(operation.get("heading", ""), "heading"), level=1)
            for block in operation.get("blocks", []):
                if block.get("type") == "paragraph":
                    document.add_paragraph(bounded_text(block.get("text", ""), "paragraph text"))
                elif block.get("type") == "bullets":
                    for item in block.get("items", []):
                        document.add_paragraph(bounded_text(item, "bullet"), style="List Bullet")
                else:
                    raise WorkbenchError("append_section supports paragraph and bullets blocks")
            changes += 1
        else:
            raise WorkbenchError(f"unsupported DOCX edit operation: {op!r}")
    output, temporary = atomic_target(output)
    try:
        document.save(temporary)
        atomic_commit(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return changes


def edit_xlsx(path: Path, spec: dict[str, Any], output: Path) -> int:
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=False)
    changes = 0
    try:
        for operation in spec["operations"]:
            op = operation.get("op")
            if op == "set_cell":
                sheet = bounded_text(operation.get("sheet"), "sheet")
                if sheet not in workbook.sheetnames:
                    raise WorkbenchError(f"sheet not found: {sheet}")
                cell = bounded_text(operation.get("cell"), "cell").upper()
                if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", cell):
                    raise WorkbenchError(f"invalid A1 cell reference: {cell}")
                workbook[sheet][cell] = safe_cell_value(operation.get("value"), bool(operation.get("allowFormula", False)))
                changes += 1
            elif op == "append_rows":
                sheet = bounded_text(operation.get("sheet"), "sheet")
                if sheet not in workbook.sheetnames:
                    raise WorkbenchError(f"sheet not found: {sheet}")
                rows = operation.get("rows", [])
                if not isinstance(rows, list):
                    raise WorkbenchError("append_rows.rows must be an array")
                for row in rows:
                    if not isinstance(row, list):
                        raise WorkbenchError("append_rows values must be arrays")
                    workbook[sheet].append([safe_cell_value(value, bool(operation.get("allowFormula", False))) for value in row])
                    changes += 1
            elif op == "add_sheet":
                name = bounded_text(operation.get("name"), "sheet name")[:31]
                if name in workbook.sheetnames:
                    raise WorkbenchError(f"sheet already exists: {name}")
                worksheet = workbook.create_sheet(name)
                for row in operation.get("rows", []):
                    worksheet.append([safe_cell_value(value, False) for value in row])
                changes += 1
            else:
                raise WorkbenchError(f"unsupported XLSX edit operation: {op!r}")
        target, temporary = atomic_target(output)
        try:
            workbook.save(temporary)
            atomic_commit(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
    finally:
        workbook.close()
    return changes


def edit_pptx(path: Path, spec: dict[str, Any], output: Path) -> int:
    from pptx import Presentation
    presentation = Presentation(path)
    changes = 0
    for operation in spec["operations"]:
        op = operation.get("op")
        if op == "replace_text":
            find = bounded_text(operation.get("find"), "find")
            replace = bounded_text(operation.get("replace", ""), "replace")
            if not find:
                raise WorkbenchError("replace_text.find cannot be empty")
            for slide in presentation.slides:
                paragraphs = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        paragraphs.extend(shape.text_frame.paragraphs)
                changes += replace_text_in_runs(paragraphs, find, replace, int(operation.get("count", 0)))
        elif op == "append_slide":
            slide_spec = operation.get("slide", {})
            if not isinstance(slide_spec, dict):
                raise WorkbenchError("append_slide.slide must be an object")
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = bounded_text(slide_spec.get("title", ""), "slide title")
            body = next((placeholder for placeholder in slide.placeholders if getattr(placeholder, "has_text_frame", False) and placeholder != slide.shapes.title), None)
            if body:
                body.text_frame.clear()
                for index, bullet in enumerate(slide_spec.get("bullets", [])):
                    paragraph = body.text_frame.paragraphs[0] if index == 0 else body.text_frame.add_paragraph()
                    paragraph.text = bounded_text(bullet, "slide bullet")
            changes += 1
        else:
            raise WorkbenchError(f"unsupported PPTX edit operation: {op!r}")
    target, temporary = atomic_target(output)
    try:
        presentation.save(temporary)
        atomic_commit(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return changes


def edit_notebook(path: Path, spec: dict[str, Any], output: Path) -> int:
    notebook = load_notebook(path)
    cells = notebook["cells"]
    changes = 0
    for operation in spec["operations"]:
        op = operation.get("op")
        if op == "clear_outputs":
            for cell in cells:
                if cell["cell_type"] == "code":
                    cell["outputs"] = []
                    cell["execution_count"] = None
                    changes += 1
        elif op == "append_cell":
            cell_spec = operation.get("cell", {})
            cell_type = cell_spec.get("type")
            source = bounded_text(cell_spec.get("source", ""), "cell source")
            if cell_type not in {"markdown", "code", "raw"}:
                raise WorkbenchError(f"unsupported notebook cell type: {cell_type!r}")
            if len(cells) >= MAX_NOTEBOOK_CELLS:
                raise WorkbenchError(f"notebook cell count exceeds {MAX_NOTEBOOK_CELLS}")
            cells.append(notebook_cell(cell_type, source, [], next_notebook_cell_id(cells)))
            changes += 1
        elif op == "replace_cell":
            index = int(operation.get("index", -1))
            if index < 0 or index >= len(cells):
                raise WorkbenchError(f"notebook cell index out of range: {index}")
            cells[index]["source"] = bounded_text(operation.get("source", ""), "cell source")
            changes += 1
        else:
            raise WorkbenchError(f"unsupported IPYNB edit operation: {op!r}")
    write_notebook(notebook, output)
    return changes


def edit_pdf(path: Path, spec: dict[str, Any], output: Path) -> int:
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader(path)
    writer = PdfWriter(clone_from=path)
    changes = 0
    for operation in spec["operations"]:
        op = operation.get("op")
        if op == "rotate_pages":
            degrees = int(operation.get("degrees", 0))
            if degrees not in {90, 180, 270}:
                raise WorkbenchError("rotate_pages.degrees must be 90, 180, or 270")
            for page_number in operation.get("pages", []):
                index = int(page_number) - 1
                if index < 0 or index >= len(writer.pages):
                    raise WorkbenchError(f"PDF page out of range: {page_number}")
                writer.pages[index].rotate(degrees)
                changes += 1
        elif op == "merge_pdf":
            for extra in operation.get("paths", []):
                extra_path = ensure_input(Path(bounded_text(extra, "merge path")))
                if extra_path.suffix.lower() != ".pdf":
                    raise WorkbenchError("merge_pdf accepts PDF paths only")
                extra_reader = PdfReader(extra_path)
                if extra_reader.is_encrypted:
                    raise WorkbenchError(f"encrypted merge source is blocked: {extra_path}")
                for page in extra_reader.pages:
                    writer.add_page(page)
                    changes += 1
        elif op == "set_metadata":
            metadata = operation.get("metadata", {})
            if not isinstance(metadata, dict):
                raise WorkbenchError("set_metadata.metadata must be an object")
            writer.add_metadata({str(key if str(key).startswith("/") else "/" + str(key)): bounded_text(str(value), "metadata value") for key, value in metadata.items()})
            changes += 1
        else:
            raise WorkbenchError(f"unsupported PDF edit operation: {op!r}")
    target, temporary = atomic_target(output)
    try:
        with temporary.open("wb") as handle:
            writer.write(handle)
        atomic_commit(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    _ = reader
    return changes


def command_edit(args: argparse.Namespace) -> dict[str, Any]:
    source = ensure_input(Path(args.input))
    output = resolve_workspace_path(Path(args.output), "output")
    if source.resolve() == output.resolve():
        raise WorkbenchError("edit output must differ from the input path")
    if output.suffix.lower() != source.suffix.lower():
        raise WorkbenchError("edit output must preserve the native file format")
    spec = read_json(Path(args.spec))
    ensure_keys(spec, {"operations"}, "edit specification")
    if not isinstance(spec.get("operations"), list) or not spec["operations"]:
        raise WorkbenchError("edit specification requires a non-empty operations array")
    editors = {".docx": edit_docx, ".xlsx": edit_xlsx, ".pptx": edit_pptx, ".pdf": edit_pdf, ".ipynb": edit_notebook}
    editor = editors.get(source.suffix.lower())
    if editor is None:
        raise WorkbenchError(f"editing is unsupported for {source.suffix.lower()}")
    changes = editor(source, spec, output)
    return {"ok": True, "command": "edit", "output": str(output.resolve()), "changes": changes, "validation": validate_file(output)}


def escape_document_html(title: str, body: str, *, table: str = "") -> str:
    return """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title><style>
:root{{color-scheme:light}} html,body{{min-height:100%;background:#fff}} body{{margin:0;color:#202020;font:15px/1.55 -apple-system,BlinkMacSystemFont,"Segoe UI",Arial,sans-serif}}
main{{max-width:980px;margin:0 auto;padding:36px}} h1,h2,h3{{line-height:1.25}} table{{border-collapse:collapse;width:100%;margin:18px 0}}
th,td{{border:1px solid #d7d7d7;padding:8px;text-align:left;vertical-align:top}} th{{background:#f3f3f3}} pre{{white-space:pre-wrap;background:#f6f6f6;padding:14px;border-radius:8px;overflow:auto}}
.slide{{min-height:460px;padding:36px;margin:0 0 24px;background:#fff;border:1px solid #e8e8e8;border-radius:10px}} .muted{{color:#666}}
</style></head><body><main><h1>{title}</h1>{body}{table}</main></body></html>""".format(title=html.escape(title), body=body, table=table)


def notebook_output_html(output: Any) -> str:
    output_type = output.get("output_type", "")
    if output_type == "stream":
        return f"<pre>{html.escape(bounded_text(output.get('text', ''), 'notebook stream output'))}</pre>"
    if output_type == "error":
        traceback = output.get("traceback", [])
        if not isinstance(traceback, list):
            traceback = []
        clean = "\n".join(re.sub(r"\x1b\[[0-9;]*m", "", bounded_text(line, "notebook traceback")) for line in traceback)
        return f"<pre>{html.escape(clean)}</pre>"
    if output_type not in {"display_data", "execute_result"}:
        return ""
    data = output.get("data", {})
    if not isinstance(data, dict):
        return ""
    rendered: list[str] = []
    image_data = data.get("image/png")
    if isinstance(image_data, str):
        try:
            decoded = base64.b64decode(image_data, validate=True)
        except (binascii.Error, ValueError):
            decoded = b""
        if decoded.startswith(b"\x89PNG\r\n\x1a\n") and len(decoded) <= MAX_NOTEBOOK_PREVIEW_BYTES:
            rendered.append(f'<img alt="Notebook output" src="data:image/png;base64,{html.escape(image_data, quote=True)}" style="max-width:100%;height:auto">')
    text_value = data.get("text/plain")
    if isinstance(text_value, list):
        text_value = "".join(str(item) for item in text_value)
    if isinstance(text_value, str):
        rendered.append(f"<pre>{html.escape(bounded_text(text_value, 'notebook text output'))}</pre>")
    return "".join(rendered)


def notebook_preview_html(path: Path) -> str:
    notebook = load_notebook(path)
    sections: list[str] = []
    preview_bytes = 0
    for index, cell in enumerate(notebook["cells"]):
        source = bounded_text(cell.get("source", ""), f"notebook cell {index + 1} source")
        preview_bytes += len(source.encode("utf-8"))
        if preview_bytes > MAX_NOTEBOOK_PREVIEW_BYTES:
            sections.append('<p class="muted">Preview truncated at the safety limit.</p>')
            break
        label = html.escape(str(cell.get("cell_type", "unknown")))
        content = f"<pre>{html.escape(source)}</pre>"
        outputs = ""
        if cell.get("cell_type") == "code":
            outputs = "".join(notebook_output_html(output) for output in cell.get("outputs", []))
        sections.append(f"<section><h2>Cell {index + 1} · {label}</h2>{content}{outputs}</section>")
    return escape_document_html(path.name, "".join(sections))


def static_html_for(path: Path) -> str:
    suffix = path.suffix.lower()
    inspected = inspect_file(path)
    if suffix == ".docx":
        data = inspected["document"]
        body = "".join(f"<p>{html.escape(text)}</p>" for text in data["paragraphs"] if text)
        tables = "".join("<table>" + "".join("<tr>" + "".join(f"<td>{html.escape(cell)}</td>" for cell in row) + "</tr>" for row in table) + "</table>" for table in data["tables"])
        return escape_document_html(data.get("title") or path.name, body, table=tables)
    if suffix == ".xlsx":
        sheets = inspected["workbook"]["sheets"]
        body = ""
        for sheet in sheets:
            body += f"<h2>{html.escape(sheet['name'])}</h2><p class=\"muted\">{sheet['rows']} rows × {sheet['columns']} columns</p><table>"
            for row_index, row in enumerate(sheet["sample"]):
                tag = "th" if row_index == 0 else "td"
                body += "<tr>" + "".join(f"<{tag}>{html.escape('' if value is None else str(value))}</{tag}>" for value in row) + "</tr>"
            body += "</table>"
        return escape_document_html(path.name, body)
    if suffix == ".pptx":
        body = "".join(f"<section class=\"slide\"><h2>Slide {slide['number']}</h2>" + "".join(f"<p>{html.escape(text)}</p>" for text in slide["texts"]) + "</section>" for slide in inspected["presentation"]["slides"])
        return escape_document_html(path.name, body)
    if suffix == ".ipynb":
        return notebook_preview_html(path)
    if suffix == ".pdf":
        pages = inspected["pdf"]["pages"]
        body = "".join(f"<section><h2>Page {page['number']}</h2><pre>{html.escape(page['text'])}</pre></section>" for page in pages)
        return escape_document_html(path.name, body)
    if suffix == ".csv":
        rows = inspected["csv"]["rows"]
        body = "<table>" + "".join("<tr>" + "".join(f"<td>{html.escape(cell)}</td>" for cell in row) + "</tr>" for row in rows) + "</table>"
        return escape_document_html(path.name, body)
    if suffix == ".html":
        source = path.read_text(encoding="utf-8")
        parser = VisibleHTMLTextParser()
        parser.feed(source)
        parser.close()
        return escape_document_html(path.name, f"<pre>{html.escape(parser.text())}</pre>")
    raise WorkbenchError(f"static HTML preview unsupported for {suffix}")


def render_pdf_pages(pdf_path: Path, output_dir: Path, dpi: int) -> list[str]:
    require_module("pypdfium2")
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(pdf_path)
    if len(document) > MAX_PDF_PAGES:
        raise WorkbenchError(f"PDF page count exceeds {MAX_PDF_PAGES}")
    scale = max(72, min(300, dpi)) / 72.0
    pages: list[str] = []
    for index in range(len(document)):
        page = document[index]
        bitmap = page.render(scale=scale, rotation=0)
        image = bitmap.to_pil()
        name = f"page-{index + 1:03d}.png"
        image.save(output_dir / name, format="PNG", optimize=True)
        pages.append(name)
        bitmap.close()
        page.close()
    document.close()
    return pages


def libreoffice_convert(source: Path, output: Path, timeout: int) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise WorkbenchError("LibreOffice is not available for this conversion")
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="document-workbench-lo-") as profile_dir, tempfile.TemporaryDirectory(prefix="document-workbench-out-") as result_dir:
        profile_uri = Path(profile_dir).resolve().as_uri()
        target_format = output.suffix.lower().lstrip(".")
        command = [soffice, "--headless", "--nologo", "--nodefault", "--nolockcheck", "--norestore", f"-env:UserInstallation={profile_uri}", "--convert-to", target_format, "--outdir", result_dir, str(source)]
        try:
            completed = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8", errors="replace", timeout=timeout, check=False, shell=False)
        except subprocess.TimeoutExpired as exc:
            raise WorkbenchError(f"LibreOffice conversion timed out after {timeout} seconds") from exc
        if completed.returncode != 0:
            raise WorkbenchError(f"LibreOffice conversion failed ({completed.returncode}): {completed.stderr.strip() or completed.stdout.strip()}")
        candidates = list(Path(result_dir).glob(f"*.{target_format}"))
        if len(candidates) != 1:
            raise WorkbenchError(f"LibreOffice did not produce exactly one .{target_format} output")
        target, temporary = atomic_target(output)
        try:
            shutil.copyfile(candidates[0], temporary)
            atomic_commit(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
    return output.resolve()


def command_convert(args: argparse.Namespace) -> dict[str, Any]:
    source = ensure_input(Path(args.input), allow_legacy=True)
    output = resolve_workspace_path(Path(args.output), "output")
    timeout = bounded_integer(args.timeout, "conversion timeout", MIN_TIMEOUT_SECONDS, MAX_CONVERSION_TIMEOUT_SECONDS)
    if source.resolve() == output.resolve():
        raise WorkbenchError("conversion output must differ from input")
    if output.suffix.lower() == ".html" and source.suffix.lower() in SUPPORTED_FORMATS:
        target, temporary = atomic_target(output)
        try:
            temporary.write_text(static_html_for(source), encoding="utf-8")
            atomic_commit(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
    elif source.suffix.lower() in OFFICE_FORMATS | LEGACY_OFFICE_FORMATS and output.suffix.lower() == ".pdf":
        libreoffice_convert(source, output, timeout)
    elif source.suffix.lower() == ".csv" and output.suffix.lower() == ".xlsx":
        require_module("openpyxl")
        rows: list[list[str]] = []
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        spec = {"title": source.stem, "sheets": [{"name": "Data", "headers": rows[0] if rows else ["Value"], "rows": rows[1:] if rows else []}]}
        create_xlsx(spec, output)
    elif source.suffix.lower() == ".xlsx" and output.suffix.lower() == ".csv":
        require_module("openpyxl")
        from openpyxl import load_workbook
        workbook = load_workbook(source, read_only=True, data_only=False)
        target, temporary = atomic_target(output)
        try:
            with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                for row in workbook.active.iter_rows(values_only=True):
                    writer.writerow(list(row))
            atomic_commit(temporary, target)
        finally:
            workbook.close()
            temporary.unlink(missing_ok=True)
    else:
        raise WorkbenchError(f"unsupported conversion: {source.suffix.lower()} -> {output.suffix.lower()}")
    return {"ok": True, "command": "convert", "input": str(source), "output": str(output.resolve()), "validation": validate_file(output)}


def command_render(args: argparse.Namespace) -> dict[str, Any]:
    source = ensure_input(Path(args.input))
    output_dir = resolve_workspace_path(Path(args.output_dir), "preview output directory")
    dpi = bounded_integer(args.dpi, "render DPI", 72, 300)
    output_dir.mkdir(parents=True, exist_ok=True)
    if any(output_dir.iterdir()) and not args.allow_nonempty:
        raise WorkbenchError("preview output directory is not empty; use a fresh directory or --allow-nonempty")
    warnings: list[str] = []
    html_name = "index.html"
    html_target, html_temporary = atomic_target(output_dir / html_name)
    try:
        html_temporary.write_text(static_html_for(source), encoding="utf-8")
        atomic_commit(html_temporary, html_target)
    finally:
        html_temporary.unlink(missing_ok=True)
    pdf_name: str | None = None
    pages: list[str] = []
    if source.suffix.lower() == ".pdf":
        pdf_target = output_dir / source.name
        if source.resolve() != pdf_target.resolve():
            copy_target, copy_temporary = atomic_target(pdf_target)
            try:
                shutil.copyfile(source, copy_temporary)
                atomic_commit(copy_temporary, copy_target)
            finally:
                copy_temporary.unlink(missing_ok=True)
        pdf_name = pdf_target.name
        pages = render_pdf_pages(pdf_target, output_dir, dpi)
    manifest = {"schemaVersion": SCHEMA_VERSION, "source": source.name, "sourceFormat": source.suffix.lower().lstrip("."), "previewKind": "office" if source.suffix.lower() in OFFICE_FORMATS else source.suffix.lower().lstrip("."), "html": html_name, "pdf": pdf_name, "pages": pages, "warnings": warnings}
    manifest_path = output_dir / "preview-manifest.json"
    manifest_target, manifest_temporary = atomic_target(manifest_path)
    try:
        manifest_temporary.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        atomic_commit(manifest_temporary, manifest_target)
    finally:
        manifest_temporary.unlink(missing_ok=True)
    return {"ok": True, "command": "render", "outputDirectory": str(output_dir), "manifest": manifest}


def command_execute_notebook(args: argparse.Namespace) -> dict[str, Any]:
    if not args.allow_execution:
        raise WorkbenchError("notebook execution requires explicit --allow-execution")
    source = ensure_input(Path(args.input))
    if source.suffix.lower() != ".ipynb":
        raise WorkbenchError("execute-notebook accepts .ipynb input only")
    output = resolve_workspace_path(Path(args.output), "output")
    timeout = bounded_integer(args.timeout, "notebook timeout", MIN_TIMEOUT_SECONDS, MAX_NOTEBOOK_TIMEOUT_SECONDS)
    if source.resolve() == output.resolve():
        raise WorkbenchError("executed notebook output must differ from input")
    require_module("nbformat")
    require_module("nbclient")
    if sys.platform == "win32" and hasattr(asyncio, "WindowsSelectorEventLoopPolicy"):
        asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())
    import nbformat
    from nbclient import NotebookClient

    notebook = nbformat.read(source, as_version=4)
    with tempfile.TemporaryDirectory(prefix="document-workbench-notebook-") as work_dir:
        client = NotebookClient(notebook, timeout=timeout, kernel_name=args.kernel or None, allow_errors=False, resources={"metadata": {"path": work_dir}})
        try:
            client.execute()
        except Exception as exc:
            raise WorkbenchError(f"notebook execution failed: {type(exc).__name__}: {exc}") from exc
    target, temporary = atomic_target(output)
    try:
        nbformat.write(notebook, temporary)
        atomic_commit(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return {"ok": True, "command": "execute-notebook", "output": str(output.resolve()), "validation": validate_file(output)}


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description=__doc__)
    subparsers = root.add_subparsers(dest="command", required=True)
    doctor = subparsers.add_parser("doctor", help="report optional dependency availability")
    doctor.set_defaults(handler=command_doctor)
    create = subparsers.add_parser("create", help="create a native document from a JSON specification")
    create.add_argument("--format", required=True, choices=["docx", "xlsx", "pptx", "pdf", "ipynb", "html"])
    create.add_argument("--spec", required=True)
    create.add_argument("--output", required=True)
    create.set_defaults(handler=command_create)
    validate = subparsers.add_parser("validate", help="validate a supported file")
    validate.add_argument("--input", required=True)
    validate.set_defaults(handler=command_validate)
    inspect = subparsers.add_parser("inspect", help="extract a bounded structural summary")
    inspect.add_argument("--input", required=True)
    inspect.add_argument("--output")
    inspect.set_defaults(handler=command_inspect)
    edit = subparsers.add_parser("edit", help="apply explicit non-destructive native edits")
    edit.add_argument("--input", required=True)
    edit.add_argument("--spec", required=True)
    edit.add_argument("--output", required=True)
    edit.set_defaults(handler=command_edit)
    convert = subparsers.add_parser("convert", help="convert between supported formats")
    convert.add_argument("--input", required=True)
    convert.add_argument("--output", required=True)
    convert.add_argument("--timeout", type=int, default=180)
    convert.set_defaults(handler=command_convert)
    render = subparsers.add_parser("render", help="create safe HTML and visual previews")
    render.add_argument("--input", required=True)
    render.add_argument("--output-dir", required=True)
    render.add_argument("--dpi", type=int, default=144)
    render.add_argument("--timeout", type=int, default=180)
    render.add_argument("--allow-nonempty", action="store_true")
    render.set_defaults(handler=command_render)
    execute = subparsers.add_parser("execute-notebook", help="execute a notebook only with explicit authorization")
    execute.add_argument("--input", required=True)
    execute.add_argument("--output", required=True)
    execute.add_argument("--allow-execution", action="store_true")
    execute.add_argument("--timeout", type=int, default=120)
    execute.add_argument("--kernel")
    execute.set_defaults(handler=command_execute_notebook)
    return root


def main() -> int:
    args = parser().parse_args()
    try:
        result = args.handler(args)
    except WorkbenchError as exc:
        emit({"ok": False, "command": getattr(args, "command", None), "error": str(exc)})
        return 2
    except KeyboardInterrupt:
        emit({"ok": False, "command": getattr(args, "command", None), "error": "cancelled"})
        return 130
    emit(result)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
