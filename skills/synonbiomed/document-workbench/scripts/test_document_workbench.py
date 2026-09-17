#!/usr/bin/env python3
"""Behavioral regression suite for the Document Workbench."""

from __future__ import annotations

import argparse
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path


sys.dont_write_bytecode = True


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
MODULE_PATH = SCRIPT_DIR / "document_workbench.py"
SPEC = importlib.util.spec_from_file_location("document_workbench", MODULE_PATH)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError(f"cannot import {MODULE_PATH}")
workbench = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(workbench)


def namespace(**values: object) -> argparse.Namespace:
    return argparse.Namespace(**values)


def write_json(path: Path, payload: dict[str, object]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


class DocumentWorkbenchTests(unittest.TestCase):
    maxDiff = None

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory(prefix="document-workbench-test-")
        self.root = Path(self.temp.name)
        self.outside = tempfile.TemporaryDirectory(prefix="document-workbench-outside-")
        self.previous_cwd = Path.cwd()
        os.chdir(self.root)

    def tearDown(self) -> None:
        os.chdir(self.previous_cwd)
        self.outside.cleanup()
        self.temp.cleanup()

    def create(self, format_name: str, example_name: str) -> Path:
        output = self.root / f"sample.{format_name}"
        spec_path = self.root / f"create-{format_name}-{example_name}"
        spec_path.write_bytes((SKILL_DIR / "examples" / example_name).read_bytes())
        result = workbench.command_create(
            namespace(format=format_name, spec=str(spec_path), output=str(output))
        )
        self.assertTrue(result["ok"])
        self.assertTrue(output.is_file())
        self.assertGreater(output.stat().st_size, 100)
        return output

    def test_cli_is_callable_from_project_workspace_outside_skill_directory(self) -> None:
        completed = subprocess.run(
            [sys.executable, str(MODULE_PATH), "doctor"],
            cwd=self.root,
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        result = json.loads(completed.stdout)
        self.assertEqual(result["command"], "doctor")
        self.assertTrue(result["ok"])

    def render(self, source: Path) -> tuple[Path, dict[str, object]]:
        output_dir = self.root / f"{source.stem}-{source.suffix.lower().lstrip('.')}-preview"
        result = workbench.command_render(
            namespace(input=str(source), output_dir=str(output_dir), dpi=96, timeout=180, allow_nonempty=False)
        )
        manifest = result["manifest"]
        self.assertTrue((output_dir / str(manifest["html"])).is_file())
        self.assertTrue((output_dir / "preview-manifest.json").is_file())
        for page in manifest["pages"]:
            self.assertTrue((output_dir / page).is_file())
        if manifest["pdf"]:
            self.assertTrue((output_dir / str(manifest["pdf"])).is_file())
        return output_dir, manifest

    def test_create_validate_inspect_edit_and_render_all_native_formats(self) -> None:
        cases = {
            "docx": "report.json",
            "xlsx": "workbook.json",
            "pptx": "deck.json",
            "pdf": "report.json",
            "ipynb": "notebook.json",
            "html": "report.json",
        }
        created: dict[str, Path] = {}
        for format_name, example in cases.items():
            with self.subTest(format=format_name):
                source = self.create(format_name, example)
                created[format_name] = source
                validation = workbench.validate_file(source)
                inspection = workbench.inspect_file(source)
                self.assertTrue(validation["valid"])
                self.assertEqual(inspection["validation"]["format"], format_name)
                preview_dir, manifest = self.render(source)
                preview_html = (preview_dir / "index.html").read_text(encoding="utf-8")
                self.assertNotRegex(preview_html, r"(?i)<script\b|<iframe\b|https?://")
                self.assertIn(":root{color-scheme:light}", preview_html)
                self.assertIn("html,body{min-height:100%;background:#fff}", preview_html)
                self.assertNotIn("prefers-color-scheme:dark", preview_html)
                self.assertNotIn("box-shadow", preview_html)
                if format_name == "pdf":
                    self.assertGreaterEqual(len(manifest["pages"]), 1)

        edit_specs: dict[str, dict[str, object]] = {
            "docx": {"operations": [{"op": "replace_text", "find": "Synon Biomed", "replace": "Validated Workbench", "count": 1}]},
            "xlsx": {"operations": [{"op": "set_cell", "sheet": "Results", "cell": "B2", "value": 41.25}]},
            "pptx": {"operations": [{"op": "replace_text", "find": "Next step", "replace": "Validated next step", "count": 1}]},
            "pdf": {"operations": [{"op": "set_metadata", "metadata": {"Subject": "Validated workbench"}}]},
            "ipynb": {"operations": [{"op": "append_cell", "cell": {"type": "markdown", "source": "## Validation complete"}}]},
        }
        for format_name, spec in edit_specs.items():
            with self.subTest(edit=format_name):
                spec_path = self.root / f"edit-{format_name}.json"
                write_json(spec_path, spec)
                revised = self.root / f"sample-revised.{format_name}"
                result = workbench.command_edit(
                    namespace(input=str(created[format_name]), spec=str(spec_path), output=str(revised))
                )
                self.assertTrue(result["ok"])
                self.assertGreaterEqual(result["changes"], 1)
                self.assertTrue(revised.is_file())

    def test_csv_round_trip_and_html_preview_are_safe(self) -> None:
        csv_source = self.root / "source.csv"
        csv_source.write_text("compound,value\nA-01,42.5\n", encoding="utf-8")
        workbook = self.root / "source.xlsx"
        workbench.command_convert(namespace(input=str(csv_source), output=str(workbook), timeout=180))
        round_trip = self.root / "round-trip.csv"
        workbench.command_convert(namespace(input=str(workbook), output=str(round_trip), timeout=180))
        self.assertIn("A-01", round_trip.read_text(encoding="utf-8-sig"))

        safe_html = self.root / "safe.html"
        safe_html.write_text("<!doctype html><h1>Visible report</h1><p>Only text is retained.</p>", encoding="utf-8")
        preview_dir, _ = self.render(safe_html)
        preview = (preview_dir / "index.html").read_text(encoding="utf-8")
        self.assertIn("Visible report", preview)
        self.assertNotIn("<!doctype html><h1>", preview)

    def test_presentation_image_preserves_source_aspect_ratio(self) -> None:
        if not all(workbench.package_status().get(name, False) for name in ("python-pptx", "Pillow")):
            self.skipTest("managed presentation image dependencies are unavailable")
        from PIL import Image
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        image_path = self.root / "wide.png"
        Image.new("RGB", (320, 120), "white").save(image_path)
        spec_path = self.root / "wide-deck.json"
        write_json(
            spec_path,
            {
                "title": "Aspect ratio check",
                "author": "Synon Biomed",
                "slides": [
                    {
                        "layout": "image",
                        "title": "Native image dimensions",
                        "image": str(image_path),
                        "caption": "The image must not be stretched.",
                    }
                ],
            },
        )
        output = self.root / "wide.pptx"
        workbench.command_create(namespace(format="pptx", spec=str(spec_path), output=str(output)))
        presentation = Presentation(output)
        pictures = [shape for shape in presentation.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pictures), 1)
        self.assertAlmostEqual(pictures[0].width / pictures[0].height, 320 / 120, places=2)

    def test_rejects_macro_active_html_and_unsafe_packages(self) -> None:
        macro = self.root / "macro.docm"
        macro.write_bytes(b"not a document")
        with self.assertRaisesRegex(workbench.WorkbenchError, "macro-enabled"):
            workbench.ensure_input(macro)

        active = self.root / "active.html"
        active.write_text('<a href="https://example.invalid">external</a>', encoding="utf-8")
        with self.assertRaisesRegex(workbench.WorkbenchError, "active or externally loaded"):
            workbench.validate_file(active)

        unsafe = self.root / "unsafe.docx"
        with zipfile.ZipFile(unsafe, "w") as archive:
            archive.writestr("../escape.xml", "unsafe")
        with self.assertRaisesRegex(workbench.WorkbenchError, "unsafe package entry path"):
            workbench.validate_file(unsafe)

        self.assertEqual(workbench.safe_cell_value("=2+2", False), "'=2+2")
        self.assertEqual(workbench.safe_cell_value("+cmd|' /C calc'!A0", False), "'+cmd|' /C calc'!A0")
        self.assertEqual(workbench.safe_cell_value("@SUM(A1:A2)", False), "'@SUM(A1:A2)")
        self.assertEqual(workbench.safe_cell_value("=2+2", True), "=2+2")

    def test_rejects_unbounded_timeouts_and_render_dpi(self) -> None:
        source = self.create("ipynb", "notebook.json")
        with self.assertRaisesRegex(workbench.WorkbenchError, "notebook timeout must be between"):
            workbench.command_execute_notebook(
                namespace(input=str(source), output=str(self.root / "executed.ipynb"), allow_execution=True, timeout=0, kernel=None)
            )

        pdf_source = self.create("pdf", "report.json")
        with self.assertRaisesRegex(workbench.WorkbenchError, "render DPI must be between"):
            workbench.command_render(
                namespace(input=str(pdf_source), output_dir=str(self.root / "bad-render"), dpi=9999, timeout=180, allow_nonempty=False)
            )

    def test_rejects_paths_outside_workspace_and_symlink_escape(self) -> None:
        outside = Path(self.outside.name)
        outside_input = outside / "outside.html"
        outside_input.write_text("<!doctype html><p>outside</p>", encoding="utf-8")
        with self.assertRaisesRegex(workbench.WorkbenchError, "inside the current project workspace"):
            workbench.command_validate(namespace(input=str(outside_input)))

        spec = self.root / "report.json"
        write_json(spec, {"title": "Workspace boundary", "blocks": [{"type": "paragraph", "text": "bounded"}]})
        with self.assertRaisesRegex(workbench.WorkbenchError, "inside the current project workspace"):
            workbench.command_create(namespace(format="html", spec=str(spec), output="../escape.html"))

        link = self.root / "outside-link"
        try:
            link.symlink_to(outside, target_is_directory=True)
        except (NotImplementedError, OSError) as exc:
            self.skipTest(f"directory symlinks are unavailable: {exc}")
        with self.assertRaisesRegex(workbench.WorkbenchError, "inside the current project workspace"):
            workbench.command_create(namespace(format="html", spec=str(spec), output=str(link / "escape.html")))

    def test_notebook_execution_requires_explicit_authorization(self) -> None:
        source = self.create("ipynb", "notebook.json")
        with self.assertRaisesRegex(workbench.WorkbenchError, "requires explicit"):
            workbench.command_execute_notebook(
                namespace(input=str(source), output=str(self.root / "executed.ipynb"), allow_execution=False, timeout=60, kernel=None)
            )

    def test_notebook_execution_uses_a_separate_output(self) -> None:
        if not all(workbench.package_status().get(name, False) for name in ("nbformat", "nbclient", "ipykernel")):
            self.skipTest("managed notebook execution dependencies are unavailable")
        source = self.create("ipynb", "notebook.json")
        output = self.root / "executed.ipynb"
        result = workbench.command_execute_notebook(
            namespace(input=str(source), output=str(output), allow_execution=True, timeout=60, kernel="python3")
        )
        self.assertTrue(result["ok"])
        inspected = workbench.inspect_file(output)
        self.assertEqual(inspected["notebook"]["cells"][1]["outputs"], 1)
        self.assertEqual(workbench.inspect_file(source)["notebook"]["cells"][1]["outputs"], 0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
